"""Turn uploads into text pages and elements. File types are detected by content, not extension."""
from __future__ import annotations

import io
import re
import zipfile

from .categories import category_for
from .errors import AppError

PNG = b"\x89PNG\r\n\x1a\n"
JPEG = b"\xff\xd8\xff"


def sniff(data: bytes) -> str:
    """Return pdf | docx | png | jpeg | other."""
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:8] == PNG:
        return "png"
    if data[:3] == JPEG:
        return "jpeg"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    return "other"


def is_image(data: bytes) -> bool:
    return sniff(data) in {"png", "jpeg"}


def extract_pages(data: bytes, max_bytes: int) -> list[tuple[int, str]]:
    """Return [(page_number, text)]. Raises AppError for oversize or unsupported files."""
    if len(data) > max_bytes:
        raise AppError(413, "File too large")
    kind = sniff(data)
    if kind == "pdf":
        try:
            import fitz  # PyMuPDF
            with fitz.open(stream=data, filetype="pdf") as doc:
                return [(i + 1, page.get_text()) for i, page in enumerate(doc)]
        except ImportError:
            raise AppError(501, "PDF support needs PyMuPDF: pip install PyMuPDF")
        except Exception:
            raise AppError(422, "This PDF could not be read")
    if kind == "docx":
        try:
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
            return [(1, "\n".join(parts))]
        except ImportError:
            raise AppError(501, "DOCX support needs python-docx: pip install python-docx")
        except Exception:
            raise AppError(422, "This Word document could not be read")
    if kind in {"png", "jpeg"}:
        # An image can still be stored and reviewed (human review / consented vision) when OCR is unavailable.
        try:
            import pytesseract
            from PIL import Image
            return [(1, pytesseract.image_to_string(Image.open(io.BytesIO(data))))]
        except Exception:
            return [(1, "")]
    try:
        return [(1, data.decode("utf-8-sig"))]
    except UnicodeDecodeError:
        raise AppError(415, "Unsupported or binary file type")


def chunk_text(text: str, size: int = 600) -> list[str]:
    """Split text into ~size-character chunks on paragraph boundaries; oversized paragraphs split by sentence."""
    pieces: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= size:
            pieces.append(para)
            continue
        current = ""
        for sentence in re.split(r"(?<=[.!?])\s+", para):
            if current and len(current) + len(sentence) + 1 > size:
                pieces.append(current)
                current = ""
            current = f"{current} {sentence}".strip()
        if current:
            pieces.append(current)
    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) > size:
            chunks.append(current)
            current = ""
        current = f"{current}\n\n{piece}".strip()
    if current:
        chunks.append(current)
    return chunks


def split_elements(text: str, page: int = 1) -> list[dict]:
    """One element per non-empty line (the whole text if it has none)."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines and text.strip():
        lines = [text.strip()]
    return [{"element_type": "text", "text": ln, "description": "", "categories": category_for(ln),
             "page": page, "bbox": None, "confidence": None, "method": "text_extraction"} for ln in lines]
